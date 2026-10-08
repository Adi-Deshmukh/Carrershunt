from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_resume_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported resume format. Use PDF, DOCX, or TXT.")

    if not content:
        raise ValueError("Resume file is empty")

    if suffix == ".txt":
        text = content.decode("utf-8", errors="replace")
    elif suffix == ".docx":
        document = Document(BytesIO(content))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
        text = "\n".join(parts)
    else:
        reader = PdfReader(BytesIO(content))
        text = "\n".join((page.extract_text() or "").strip() for page in reader.pages)

    normalized = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(normalized) < 50:
        raise ValueError("Could not extract enough text from the resume.")
    return normalized
