import copy
import json
import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from app.db.models import CandidateProfile, Job
from app.schemas.agent import ResumePlan
from app.services.ai_service import AIService
from app.services.matching import match_job

OUTPUT_DIR = Path("generated_resumes")
OUTPUT_DIR.mkdir(exist_ok=True)

MASTER_DIR = OUTPUT_DIR / "masters"
MASTER_DIR.mkdir(exist_ok=True)


def master_resume_path(candidate_id: int) -> Path:
    return MASTER_DIR / f"candidate_{candidate_id}_master.docx"


def _candidate_dict(candidate: CandidateProfile) -> dict:
    return {
        "name": candidate.name,
        "education": candidate.education,
        "graduation_year": candidate.graduation_year,
        "experience_years": candidate.experience_years,
        "location": candidate.location,
        "work_authorization": candidate.work_authorization,
        "resume_text": candidate.resume_text,
        "evidence": json.loads(candidate.evidence_json or "{}"),
    }


def _job_dict(job: Job) -> dict:
    return {
        "title": job.title,
        "location": job.location,
        "employment_type": job.employment_type,
        "department": job.department,
        "description": job.description,
        "job_url": job.job_url,
        "company": job.company.name if job.company else "",
    }


def _norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (value or "").lower()).strip()


def _paragraphs(doc: Document) -> list[Paragraph]:
    return list(doc.paragraphs)


def _is_heading(p: Paragraph) -> bool:
    style = (p.style.name if p.style else "").lower()
    return "heading" in style or style in {"title", "subtitle"}


def _find_heading(doc: Document, aliases: set[str]) -> int | None:
    aliases = {_norm(a) for a in aliases}
    for i, p in enumerate(_paragraphs(doc)):
        if _norm(p.text) in aliases or (_is_heading(p) and any(a in _norm(p.text) for a in aliases)):
            return i
    return None


def _replace_text_preserving_runs(paragraph: Paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _clone_paragraph_after(paragraph: Paragraph, source: Paragraph | None = None) -> Paragraph:
    source = source or paragraph
    new_p = copy.deepcopy(source._p)
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def _replace_section(doc: Document, aliases: set[str], paragraphs: list[str]) -> bool:
    idx = _find_heading(doc, aliases)
    if idx is None:
        return False

    ps = _paragraphs(doc)
    heading = ps[idx]
    end = len(ps)
    for j in range(idx + 1, len(ps)):
        if _is_heading(ps[j]):
            end = j
            break

    existing = ps[idx + 1:end]
    # Keep the first body paragraph as the formatting template.
    template = existing[0] if existing else None
    for p in reversed(existing):
        p._element.getparent().remove(p._element)

    anchor = heading
    for text in paragraphs:
        if not text.strip():
            continue
        new_p = _clone_paragraph_after(anchor, template)
        _replace_text_preserving_runs(new_p, text)
        anchor = new_p
    return True


def _append_section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for text in paragraphs:
        doc.add_paragraph(text)


def _write_with_master_template(candidate: CandidateProfile, job: Job, content: dict, path: Path) -> bool:
    master = master_resume_path(candidate.id)
    if not master.exists():
        return False

    doc = Document(master)
    summary = [content.get("summary", "").strip()] if content.get("summary") else []
    skills = [", ".join(str(x) for x in content.get("skills", []) if str(x).strip())]
    skills = [x for x in skills if x]
    projects = []
    for project in content.get("projects", []):
        name = str(project.get("name", "")).strip()
        if name:
            projects.append(name)
        projects.extend(f"• {str(b).strip()}" for b in project.get("bullets", []) if str(b).strip())

    _replace_section(doc, {"summary", "professional summary", "profile", "objective"}, summary)
    _replace_section(doc, {"skills", "technical skills", "core skills", "technologies"}, skills)
    if projects:
        if not _replace_section(doc, {"projects", "relevant projects", "selected projects"}, projects):
            _append_section(doc, "Relevant Projects", projects)

    # Preserve the original document's layout, styles, headers, footers, tables,
    # margins and typography. Only targeted section content is changed.
    doc.save(path)
    return True


def _write_docx(candidate: CandidateProfile, job: Job, content: dict, suffix: str) -> Path:
    path = OUTPUT_DIR / f"candidate_{candidate.id}_job_{job.id}{suffix}.docx"
    if not _write_with_master_template(candidate, job, content, path):
        doc = Document()
        doc.add_heading(candidate.name, level=0)
        doc.add_paragraph(content.get("summary", ""))
        if content.get("skills"):
            doc.add_heading("Skills", level=1)
            doc.add_paragraph(", ".join(str(x) for x in content["skills"]))
        if content.get("projects"):
            doc.add_heading("Relevant Projects", level=1)
            for project in content["projects"]:
                doc.add_heading(project.get("name", "Project"), level=2)
                for bullet in project.get("bullets", []):
                    doc.add_paragraph(bullet, style="List Bullet")
        doc.save(path)
    return path


def generate_docx_from_plan(candidate: CandidateProfile, job: Job, plan: ResumePlan) -> tuple[Path, dict]:
    content = plan.model_dump()
    path = _write_docx(candidate, job, content, "_tailored")
    return path, content
def generate_docx(candidate: CandidateProfile, job: Job) -> tuple[Path, dict]:
    match = match_job(job, candidate).model_dump()
    content = AIService().tailor_resume(_candidate_dict(candidate), _job_dict(job), match)
    path = _write_docx(candidate, job, content, "")
    return path, content
